"""007 · Generación del Excel de la cotización con el theme Capsulab.

Lógica **pura** (sin FastAPI): recibe las líneas de la cotización y devuelve los
BYTES de un `.xlsx`. El endpoint (`rutas/solicitudes.py`) mapea la propuesta
persistida (spec 004) a `LineaCotizacion` y sirve estos bytes como descarga.

Dos vistas, mismo theme (archivo modelo `PF - Fuchs cerros.xlsx`):

* **interno** (default): el documento de trabajo de la agencia. Banda de título
  negra (Arial 25 blanca), encabezado verde `FF8ED873` con las 9 columnas
  `Item · PROVEEDOR · DESCRIPCIÓN · Cantidad · días · valor unitario · COSTO ·
  MARGEN · VALOR FINAL`. `COSTO = Cantidad×días×valor`, `VALOR FINAL =
  COSTO/(1−margen)`, subtotal amarillo y valor venta cian.
* **cliente**: lo que se le manda a la marca. Solo precios de **venta**; NO escribe
  costo, margen ni proveedor en ninguna celda (no se filtran aunque "desoculten"
  columnas). Columnas `Item · DESCRIPCIÓN · Cantidad · VALOR`, con
  `VALOR = Cantidad×días×valor/(1−margen)` y un total cian.

`valor_unitario` es el **costo** unitario (lo que Javo trae del Drive); el margen
se aplica para llegar al precio de venta.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# --- Theme Capsulab (ARGB) ---------------------------------------------------
NEGRO = "FF000000"
BLANCO = "FFFFFFFF"
VERDE = "FF8ED873"
AMARILLO = "FFFFFF00"
CIAN = "FF00FFFF"
FUENTE = "Arial"

# Formatos de número (extraídos del archivo modelo).
FMT_CLP = '_ "$"* #,##0_ ;_ "$"* \\-#,##0_ ;_ "$"* "-"_ ;_ @_ '
FMT_FINAL = '#,##0_ ;[Red]\\-#,##0\\ '
FMT_PCT = "0%"

# Columnas de la vista INTERNA (A..I) y sus anchos.
COLUMNAS = [
    ("Item", 26),
    ("PROVEEDOR", 22),
    ("DESCRIPCIÓN", 52),
    ("Cantidad", 10),
    ("días", 8),
    ("valor unitario", 16),
    ("COSTO", 16),
    ("MARGEN", 11),
    ("VALOR FINAL", 18),
]
N_COLS = len(COLUMNAS)

# Columnas de la vista CLIENTE (A..D): sin costo, margen ni proveedor.
COLUMNAS_CLIENTE = [
    ("Item", 30),
    ("DESCRIPCIÓN", 56),
    ("Cantidad", 12),
    ("VALOR", 18),
]

FILA_TITULO = 1
FILA_ENCABEZADO = 2
FILA_DATOS = 3


@dataclass
class LineaCotizacion:
    """Una línea de la cotización. `valor_unitario` es el COSTO unitario."""

    item: str
    descripcion: str = ""
    proveedor: str = ""
    cantidad: int = 1
    dias: int = 1
    valor_unitario: float = 0.0


def _relleno(color: str) -> PatternFill:
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


def _anchos(ws: Worksheet, columnas: list[tuple[str, int]]) -> None:
    for i, (_, ancho) in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _banda_titulo(ws: Worksheet, titulo: str, n_cols: int) -> None:
    """Banda superior negra con Arial 25 blanca, fusionada a lo ancho de la tabla."""
    ws.merge_cells(
        start_row=FILA_TITULO, start_column=1, end_row=FILA_TITULO, end_column=n_cols
    )
    celda = ws.cell(row=FILA_TITULO, column=1, value=titulo)
    celda.fill = _relleno(NEGRO)
    celda.font = Font(name=FUENTE, size=25, bold=True, color=BLANCO)
    celda.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[FILA_TITULO].height = 34


def _encabezado(ws: Worksheet, columnas: list[tuple[str, int]]) -> None:
    """Fila de encabezado verde (Arial bold)."""
    verde = _relleno(VERDE)
    for c, (etiqueta, _) in enumerate(columnas, start=1):
        celda = ws.cell(row=FILA_ENCABEZADO, column=c, value=etiqueta)
        celda.fill = verde
        celda.font = Font(name=FUENTE, size=11, bold=True)
        celda.alignment = Alignment(vertical="center", wrap_text=True)


def generar_excel_cotizacion(
    lineas: list[LineaCotizacion],
    *,
    titulo: str = "Cotización",
    margen: float = 0.40,
    vista: Literal["interno", "cliente"] = "interno",
) -> bytes:
    """Arma el libro de la cotización y devuelve sus bytes (.xlsx).

    `margen` es la fracción de margen (0.40 = 40 %). `vista`:
      * `"interno"` (default): costos + margen + valor final (doc de la agencia).
      * `"cliente"`: solo precios de venta (lo que se le manda a la marca).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Cotización"
    if vista == "cliente":
        _construir_cliente(ws, lineas, titulo=titulo, margen=margen)
    else:
        _construir_interno(ws, lineas, titulo=titulo, margen=margen)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _construir_interno(
    ws: Worksheet, lineas: list[LineaCotizacion], *, titulo: str, margen: float
) -> None:
    _anchos(ws, COLUMNAS)
    _banda_titulo(ws, titulo, N_COLS)
    _encabezado(ws, COLUMNAS)

    divisor = f"{1 - margen:g}"  # 0.40 → "0.6"; limpia ruido de float
    fila_ini = FILA_DATOS
    fuente = Font(name=FUENTE, size=11)
    for n, linea in enumerate(lineas):
        r = fila_ini + n
        ws.cell(row=r, column=1, value=linea.item).font = fuente
        ws.cell(row=r, column=2, value=linea.proveedor).font = fuente
        celda_desc = ws.cell(row=r, column=3, value=linea.descripcion)
        celda_desc.font = fuente
        celda_desc.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=4, value=linea.cantidad).font = fuente
        ws.cell(row=r, column=5, value=linea.dias).font = fuente
        f = ws.cell(row=r, column=6, value=linea.valor_unitario)
        f.number_format = FMT_CLP
        f.font = fuente
        g = ws.cell(row=r, column=7, value=f"=D{r}*E{r}*F{r}")
        g.number_format = FMT_CLP
        g.font = fuente
        h = ws.cell(row=r, column=8, value=f"=1-G{r}/I{r}")
        h.number_format = FMT_PCT
        h.font = fuente
        i = ws.cell(row=r, column=9, value=f"=G{r}/{divisor}")
        i.number_format = FMT_FINAL
        i.font = Font(name=FUENTE, size=11, bold=True)

    fila_fin = fila_ini + len(lineas) - 1 if lineas else fila_ini

    # Subtotal (amarillo).
    fila_sub = fila_fin + 1
    amarillo = _relleno(AMARILLO)
    et_sub = ws.cell(row=fila_sub, column=6, value="SUB TOTAL COSTOS")
    et_sub.font = Font(name=FUENTE, size=11, bold=True)
    sub_g = ws.cell(row=fila_sub, column=7, value=f"=SUM(G{fila_ini}:G{fila_fin})")
    sub_g.number_format = FMT_CLP
    sub_g.font = Font(name=FUENTE, size=11, bold=True)
    sub_h = ws.cell(row=fila_sub, column=8, value=f"=1-G{fila_sub}/I{fila_sub}")
    sub_h.number_format = FMT_PCT
    sub_i = ws.cell(row=fila_sub, column=9, value=f"=SUM(I{fila_ini}:I{fila_fin})")
    sub_i.number_format = FMT_FINAL
    sub_i.font = Font(name=FUENTE, size=11, bold=True)
    for c in range(6, N_COLS + 1):
        ws.cell(row=fila_sub, column=c).fill = amarillo

    # Valor venta (cian).
    fila_venta = fila_sub + 1
    cian = _relleno(CIAN)
    et_venta = ws.cell(row=fila_venta, column=8, value="VALOR VENTA")
    et_venta.font = Font(name=FUENTE, size=11, bold=True)
    venta_i = ws.cell(row=fila_venta, column=9, value=f"=I{fila_sub}")
    venta_i.number_format = FMT_CLP
    venta_i.font = Font(name=FUENTE, size=11, bold=True)
    for c in range(8, N_COLS + 1):
        ws.cell(row=fila_venta, column=c).fill = cian


def _construir_cliente(
    ws: Worksheet, lineas: list[LineaCotizacion], *, titulo: str, margen: float
) -> None:
    """Vista para el cliente: solo precios de venta. NUNCA escribe costo, margen ni
    proveedor en una celda (no se filtran aunque desoculten columnas)."""
    n_cols = len(COLUMNAS_CLIENTE)
    _anchos(ws, COLUMNAS_CLIENTE)
    _banda_titulo(ws, titulo, n_cols)
    _encabezado(ws, COLUMNAS_CLIENTE)

    fila_ini = FILA_DATOS
    fuente = Font(name=FUENTE, size=11)
    for n, linea in enumerate(lineas):
        r = fila_ini + n
        ws.cell(row=r, column=1, value=linea.item).font = fuente
        celda_desc = ws.cell(row=r, column=2, value=linea.descripcion)
        celda_desc.font = fuente
        celda_desc.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=3, value=linea.cantidad).font = fuente
        # Precio de VENTA de la línea (ya con margen). Se escribe el número calculado,
        # NO una fórmula que referencie el costo: así el costo no queda en el archivo.
        venta = round(linea.cantidad * linea.dias * linea.valor_unitario / (1 - margen))
        d = ws.cell(row=r, column=4, value=venta)
        d.number_format = FMT_CLP
        d.font = Font(name=FUENTE, size=11, bold=True)

    fila_fin = fila_ini + len(lineas) - 1 if lineas else fila_ini

    # Total (cian).
    fila_total = fila_fin + 1
    cian = _relleno(CIAN)
    et = ws.cell(row=fila_total, column=3, value="VALOR TOTAL")
    et.font = Font(name=FUENTE, size=11, bold=True)
    et.alignment = Alignment(horizontal="right")
    total = ws.cell(row=fila_total, column=4, value=f"=SUM(D{fila_ini}:D{fila_fin})")
    total.number_format = FMT_CLP
    total.font = Font(name=FUENTE, size=11, bold=True)
    for c in range(3, n_cols + 1):
        ws.cell(row=fila_total, column=c).fill = cian
