"""007 · Tests del generador de Excel de la cotización (lógica pura, sin FastAPI).

Se prueba abriendo los BYTES generados con openpyxl (sin motor de cálculo externo):
se asertan valores, fórmulas, fills y fuentes del theme Capsulab (archivo modelo
`PF - Fuchs cerros.xlsx`): banda de título negra/blanca, encabezado verde, COSTO y
VALOR FINAL como fórmulas, subtotal amarillo y valor venta cian.
"""
from io import BytesIO

import openpyxl

from app.servicios.exportador_excel import (
    LineaCotizacion,
    generar_excel_cotizacion,
)

VERDE = "FF8ED873"
AMARILLO = "FFFFFF00"
CIAN = "FF00FFFF"
NEGRO = "FF000000"
BLANCO = "FFFFFFFF"

_LINEAS = [
    LineaCotizacion(
        item="Promotoras uniformadas",
        descripcion="6h/día · 3 tiendas",
        proveedor="ISABEL",
        cantidad=6,
        dias=4,
        valor_unitario=240000,
    ),
    LineaCotizacion(
        item="Catering",
        descripcion="Té + bocados",
        proveedor="Muzia",
        cantidad=1,
        dias=2,
        valor_unitario=380000,
    ),
]


def _hoja(lineas=None, **kw):
    datos = _LINEAS if lineas is None else lineas
    bytes_ = generar_excel_cotizacion(datos, **kw)
    wb = openpyxl.load_workbook(BytesIO(bytes_))
    return wb.active


def _fill(celda):
    f = celda.fill
    return f.fgColor.rgb if f and f.patternType else None


def _buscar(ws, texto):
    for fila in ws.iter_rows():
        for celda in fila:
            if isinstance(celda.value, str) and texto in celda.value:
                return celda.row, celda.column
    raise AssertionError(f"no se encontró {texto!r} en la hoja")


def test_genera_xlsx_valido_reabrible():
    bytes_ = generar_excel_cotizacion(_LINEAS)
    assert isinstance(bytes_, bytes)
    assert bytes_[:2] == b"PK"  # firma ZIP de un .xlsx
    wb = openpyxl.load_workbook(BytesIO(bytes_))
    assert wb.active is not None


def test_banda_titulo_negra_blanca():
    ws = _hoja(titulo="Activación Metro")
    # El título está en la primera celda (fusionada a lo ancho).
    titulo = ws["A1"]
    assert titulo.value == "Activación Metro"
    assert _fill(titulo) == NEGRO
    assert titulo.font.name == "Arial"
    assert titulo.font.size == 25
    assert titulo.font.bold is True
    assert titulo.font.color.rgb == BLANCO


def test_encabezado_verde_con_columnas_en_orden():
    ws = _hoja()
    fila_h, _ = _buscar(ws, "VALOR FINAL")
    etiquetas = [ws.cell(row=fila_h, column=c).value for c in range(1, 10)]
    assert etiquetas == [
        "Item",
        "PROVEEDOR",
        "DESCRIPCIÓN",
        "Cantidad",
        "días",
        "valor unitario",
        "COSTO",
        "MARGEN",
        "VALOR FINAL",
    ]
    # Verde Capsulab en toda la fila de encabezado.
    for c in range(1, 10):
        assert _fill(ws.cell(row=fila_h, column=c)) == VERDE
        assert ws.cell(row=fila_h, column=c).font.bold is True


def test_fila_datos_valores_y_formulas():
    ws = _hoja()
    fila_h, _ = _buscar(ws, "VALOR FINAL")
    r = fila_h + 1  # primera fila de datos
    assert ws.cell(row=r, column=1).value == "Promotoras uniformadas"
    assert ws.cell(row=r, column=2).value == "ISABEL"
    assert ws.cell(row=r, column=3).value == "6h/día · 3 tiendas"
    assert ws.cell(row=r, column=4).value == 6  # Cantidad
    assert ws.cell(row=r, column=5).value == 4  # días
    assert ws.cell(row=r, column=6).value == 240000  # valor unitario (costo)
    # COSTO = Cantidad × días × valor unitario (fórmula)
    assert ws.cell(row=r, column=7).value == f"=D{r}*E{r}*F{r}"
    # VALOR FINAL = COSTO / (1 - margen); margen por defecto 0.40 → /0.6
    assert ws.cell(row=r, column=9).value == f"=G{r}/0.6"
    # MARGEN derivado
    assert ws.cell(row=r, column=8).value == f"=1-G{r}/I{r}"


def test_subtotal_amarillo_y_valor_venta_cian():
    ws = _hoja()
    fila_h, _ = _buscar(ws, "VALOR FINAL")
    ini = fila_h + 1
    fin = ini + len(_LINEAS) - 1

    fila_sub, _ = _buscar(ws, "SUB TOTAL")
    assert _fill(ws.cell(row=fila_sub, column=7)) == AMARILLO
    assert ws.cell(row=fila_sub, column=7).value == f"=SUM(G{ini}:G{fin})"
    assert ws.cell(row=fila_sub, column=9).value == f"=SUM(I{ini}:I{fin})"

    fila_venta, _ = _buscar(ws, "VALOR VENTA")
    assert _fill(ws.cell(row=fila_venta, column=9)) == CIAN
    assert ws.cell(row=fila_venta, column=9).value == f"=I{fila_sub}"


def test_margen_custom_cambia_el_divisor():
    ws = _hoja(margen=0.5)  # divisor 1-0.5 = 0.5
    fila_h, _ = _buscar(ws, "VALOR FINAL")
    r = fila_h + 1
    assert ws.cell(row=r, column=9).value == f"=G{r}/0.5"


def test_una_fila_por_componente():
    ws = _hoja()
    fila_h, _ = _buscar(ws, "VALOR FINAL")
    fila_sub, _ = _buscar(ws, "SUB TOTAL")
    # filas de datos entre encabezado y subtotal == número de componentes
    assert fila_sub - (fila_h + 1) == len(_LINEAS)
