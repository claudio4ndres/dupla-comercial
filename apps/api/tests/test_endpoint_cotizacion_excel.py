"""007 · Tests del endpoint GET /solicitudes/{id}/cotizacion.xlsx.

El front descarga la cotización como .xlsx con el theme Capsulab. El endpoint reusa
el repo de propuestas (spec 004) y la empresa del JWT: sólo la cotización de la
empresa del usuario (RLS), 404 si no hay propuesta. Repo y empresa se inyectan con
dobles en memoria (cero red, cero LLM).
"""
from io import BytesIO
from uuid import uuid4

import openpyxl
from fastapi.testclient import TestClient

from app.dependencias import (
    obtener_empresa_actual,
    obtener_repositorio_propuestas,
)
from app.main import app
from app.repositorios.propuestas import (
    ComponentePropuesta,
    Propuesta,
    RepositorioPropuestasEnMemoria,
)

EMPRESA_A = uuid4()
EMPRESA_B = uuid4()
CONTENT_TYPE_XLSX = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)


def _cliente_http(repo, empresa_id=EMPRESA_A):
    app.dependency_overrides[obtener_repositorio_propuestas] = lambda: repo
    app.dependency_overrides[obtener_empresa_actual] = lambda: empresa_id
    return TestClient(app)


def teardown_function():
    app.dependency_overrides.clear()


def _propuesta(solicitud_id, empresa_id=EMPRESA_A):
    return Propuesta(
        id=uuid4(),
        empresa_id=empresa_id,
        solicitud_id=solicitud_id,
        total=2900000,
        estado="borrador",
        componentes=[
            ComponentePropuesta(
                nombre="Promotoras uniformadas",
                detalle="6h/día · 3 tiendas",
                proveedor="ISABEL",
                cantidad=6,
                valor_unitario=240000,
            ),
            ComponentePropuesta(
                nombre="Catering",
                detalle="Té + bocados",
                cantidad=1,
                valor_unitario=380000,
            ),
        ],
    )


def test_descarga_xlsx_valido():
    sol = uuid4()
    repo = RepositorioPropuestasEnMemoria([_propuesta(sol)])
    http = _cliente_http(repo)

    r = http.get(f"/solicitudes/{sol}/cotizacion.xlsx")

    assert r.status_code == 200
    assert r.headers["content-type"].startswith(CONTENT_TYPE_XLSX)
    assert "attachment" in r.headers.get("content-disposition", "")
    assert r.content[:2] == b"PK"  # firma ZIP
    # Reabrible y con los datos de la propuesta.
    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    valores = [c.value for fila in ws.iter_rows() for c in fila]
    assert "Promotoras uniformadas" in valores
    assert "VALOR FINAL" in valores  # encabezado del theme


def test_margen_query_param_cambia_el_divisor():
    sol = uuid4()
    repo = RepositorioPropuestasEnMemoria([_propuesta(sol)])
    http = _cliente_http(repo)

    r = http.get(f"/solicitudes/{sol}/cotizacion.xlsx?margen=0.5")

    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    # La primera fila de datos usa el divisor 1-0.5 = 0.5 en VALOR FINAL.
    formulas = [
        c.value
        for fila in ws.iter_rows()
        for c in fila
        if isinstance(c.value, str) and c.value.startswith("=G") and "/0.5" in c.value
    ]
    assert formulas, "se esperaba VALOR FINAL con divisor 0.5"


def test_columna_proveedor_poblada():
    # El proveedor persistido del componente baja a la columna PROVEEDOR (B) del Excel.
    sol = uuid4()
    repo = RepositorioPropuestasEnMemoria([_propuesta(sol)])
    http = _cliente_http(repo)

    r = http.get(f"/solicitudes/{sol}/cotizacion.xlsx")

    ws = openpyxl.load_workbook(BytesIO(r.content)).active
    col_b = [c.value for fila in ws.iter_rows() for c in fila if c.column == 2]
    assert "ISABEL" in col_b


def test_vista_cliente_no_expone_costos():
    # ?vista=cliente → archivo sin COSTO/MARGEN ni el costo unitario (240000).
    sol = uuid4()
    repo = RepositorioPropuestasEnMemoria([_propuesta(sol)])
    http = _cliente_http(repo)

    r = http.get(f"/solicitudes/{sol}/cotizacion.xlsx?vista=cliente")

    assert r.status_code == 200
    assert "cotizacion-cliente-" in r.headers.get("content-disposition", "")
    planos = [
        c.value
        for fila in openpyxl.load_workbook(BytesIO(r.content)).active.iter_rows()
        for c in fila
    ]
    assert "COSTO" not in planos
    assert "MARGEN" not in planos
    assert 240000 not in planos


def test_sin_propuesta_devuelve_404():
    repo = RepositorioPropuestasEnMemoria([])
    http = _cliente_http(repo)

    r = http.get(f"/solicitudes/{uuid4()}/cotizacion.xlsx")

    assert r.status_code == 404


def test_aislamiento_multi_tenant_no_descarga_de_otra_empresa():
    sol = uuid4()
    repo = RepositorioPropuestasEnMemoria([_propuesta(sol, empresa_id=EMPRESA_B)])
    http = _cliente_http(repo, empresa_id=EMPRESA_A)

    r = http.get(f"/solicitudes/{sol}/cotizacion.xlsx")

    assert r.status_code == 404
