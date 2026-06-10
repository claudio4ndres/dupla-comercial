"""T18 · Tests del endpoint GET /solicitudes/{id}/propuesta.pptx.

El front descarga la propuesta como un DECK .pptx (cara comercial, vista cliente).
El endpoint reusa el repo de propuestas (spec 004) y la empresa del JWT: sólo la
propuesta de la empresa del usuario (RLS), 404 si no hay. Repo y empresa se inyectan
con dobles en memoria (cero red, cero LLM). Se verifica que el costo NO se filtra.
"""
from io import BytesIO
from uuid import uuid4

from fastapi.testclient import TestClient
from pptx import Presentation

from app.dependencias import (
    obtener_empresa_actual,
    obtener_repositorio_propuestas,
)
from app.main import app
from app.repositorios.propuestas import (
    ComponentePropuesta,
    Propuesta,
    RepositorioPropuestasEnMemoria,
)

EMPRESA_A = uuid4()
EMPRESA_B = uuid4()
CONTENT_TYPE_PPTX = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


def _cliente_http(repo, empresa_id=EMPRESA_A):
    app.dependency_overrides[obtener_repositorio_propuestas] = lambda: repo
    app.dependency_overrides[obtener_empresa_actual] = lambda: empresa_id
    return TestClient(app)


def teardown_function():
    app.dependency_overrides.clear()


def _propuesta(solicitud_id, empresa_id=EMPRESA_A):
    return Propuesta(
        id=uuid4(),
        empresa_id=empresa_id,
        solicitud_id=solicitud_id,
        total=2900000,
        estado="borrador",
        componentes=[
            ComponentePropuesta(
                nombre="Promotoras uniformadas",
                detalle="6h/día · 3 tiendas",
                proveedor="ISABEL",
                cantidad=6,
                valor_unitario=240000,
            ),
            ComponentePropuesta(
                nombre="Catering",
                detalle="Té + bocados",
                cantidad=1,
                valor_unitario=380000,
            ),
        ],
    )


def _textos(prs) -> list[str]:
    out: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if shape.has_table:
                for fila in shape.table.rows:
                    for celda in fila.cells:
                        out.append(celda.text)
    return out


def test_descarga_pptx_valido():
    sol = uuid4()
    repo = RepositorioPropuestasEnMemoria([_propuesta(sol)])
    http = _cliente_http(repo)

    r = http.get(f"/solicitudes/{sol}/propuesta.pptx")

    assert r.status_code == 200
    assert r.headers["content-type"].startswith(CONTENT_TYPE_PPTX)
    disp = r.headers.get("content-disposition", "")
    assert "attachment" in disp
    assert f"cotizacion-{sol}.pptx" in disp
    assert r.content[:2] == b"PK"  # firma ZIP
    # Reabrible y con los datos de la propuesta.
    prs = Presentation(BytesIO(r.content))
    assert len(prs.slides) >= 3
    texto = "\n".join(_textos(prs))
    assert "Promotoras uniformadas" in texto


def test_no_expone_costo_ni_margen():
    sol = uuid4()
    repo = RepositorioPropuestasEnMemoria([_propuesta(sol)])
    http = _cliente_http(repo)

    r = http.get(f"/solicitudes/{sol}/propuesta.pptx")

    prs = Presentation(BytesIO(r.content))
    texto = "\n".join(_textos(prs))
    assert "240000" not in texto
    assert "240.000" not in texto
    assert "ISABEL" not in texto
    assert "MARGEN" not in texto.upper()


def test_sin_propuesta_devuelve_404():
    repo = RepositorioPropuestasEnMemoria([])
    http = _cliente_http(repo)

    r = http.get(f"/solicitudes/{uuid4()}/propuesta.pptx")

    assert r.status_code == 404


def test_aislamiento_multi_tenant_no_descarga_de_otra_empresa():
    sol = uuid4()
    repo = RepositorioPropuestasEnMemoria([_propuesta(sol, empresa_id=EMPRESA_B)])
    http = _cliente_http(repo, empresa_id=EMPRESA_A)

    r = http.get(f"/solicitudes/{sol}/propuesta.pptx")

    assert r.status_code == 404
