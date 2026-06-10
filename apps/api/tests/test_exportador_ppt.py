"""T18 · Tests del generador del DECK PPTX de la cotización (lógica pura, sin FastAPI).

Es la **cara comercial**: la versión cliente de la propuesta. Solo precios de VENTA;
NUNCA costos ni margen (mismo criterio que el Excel cliente que ya existe). Se prueba
abriendo los BYTES generados con python-pptx: firma ZIP, deck reabrible, ≥3 slides, el
título aparece en la portada y el COSTO (p.ej. 240000) NO se filtra a ningún texto.
"""
from io import BytesIO

from pptx import Presentation

from app.servicios.exportador_excel import LineaCotizacion
from app.servicios.exportador_ppt import generar_ppt_cotizacion

_LINEAS = [
    LineaCotizacion(
        item="Promotoras uniformadas",
        descripcion="6h/día · 3 tiendas",
        proveedor="ISABEL",
        cantidad=6,
        dias=4,
        valor_unitario=240000,  # COSTO unitario
    ),
    LineaCotizacion(
        item="Catering",
        descripcion="Té + bocados",
        proveedor="Muzia",
        cantidad=1,
        dias=2,
        valor_unitario=380000,  # COSTO unitario
    ),
]


def _textos(prs) -> list[str]:
    """Todo el texto del deck (de cualquier shape, incluidas las celdas de tabla)."""
    out: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
            if shape.has_table:
                for fila in shape.table.rows:
                    for celda in fila.cells:
                        out.append(celda.text)
    return out


def test_genera_pptx_valido_reabrible():
    bytes_ = generar_ppt_cotizacion(_LINEAS, titulo="Cotización · 212 VIP Black")
    assert isinstance(bytes_, bytes)
    assert bytes_[:2] == b"PK"  # firma ZIP de un .pptx
    prs = Presentation(BytesIO(bytes_))
    assert len(prs.slides) >= 3


def test_portada_contiene_marca_y_titulo():
    bytes_ = generar_ppt_cotizacion(_LINEAS, titulo="Cotización · 212 VIP Black")
    prs = Presentation(BytesIO(bytes_))
    portada = "\n".join(
        s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame
    )
    assert "Cotización · 212 VIP Black" in portada
    assert "Capsulab" in portada


def test_tabla_columnas_y_valor_de_venta():
    bytes_ = generar_ppt_cotizacion(_LINEAS, titulo="X", margen=0.40)
    prs = Presentation(BytesIO(bytes_))
    # La tabla está en la slide 2.
    tabla = next(s.table for s in prs.slides[1].shapes if s.has_table)
    encabezados = [c.text for c in tabla.rows[0].cells]
    assert encabezados == ["Ítem", "Descripción", "Cantidad", "VALOR"]
    # Primera línea: Promotoras (cant 6 · días 4 · costo 240000) → venta con
    # margen 0.40: 6×4×240000/0.6 = 9.600.000.
    fila = [c.text for c in tabla.rows[1].cells]
    assert fila[0] == "Promotoras uniformadas"
    assert fila[2] == "6"
    assert "9.600.000" in fila[3]


def test_valor_total_de_venta():
    bytes_ = generar_ppt_cotizacion(_LINEAS, titulo="X", margen=0.40)
    prs = Presentation(BytesIO(bytes_))
    # Promotoras 9.600.000 + Catering 1×2×380000/0.6 = 1.266.667 → total 10.866.667.
    texto = "\n".join(_textos(prs))
    assert "10.866.667" in texto


def test_no_filtra_costo_ni_margen():
    # El deck que se le manda al cliente NUNCA contiene el costo unitario ni el margen.
    bytes_ = generar_ppt_cotizacion(_LINEAS, titulo="X", margen=0.40)
    prs = Presentation(BytesIO(bytes_))
    texto = "\n".join(_textos(prs))
    assert "240000" not in texto  # costo unitario crudo
    assert "240.000" not in texto  # costo unitario formateado
    assert "MARGEN" not in texto.upper()
    assert "40%" not in texto
    assert "ISABEL" not in texto  # tampoco el proveedor


def test_margen_custom_cambia_la_venta():
    # Con margen 0.5 el divisor es 0.5: Promotoras 6×4×240000/0.5 = 11.520.000.
    bytes_ = generar_ppt_cotizacion(_LINEAS, titulo="X", margen=0.5)
    prs = Presentation(BytesIO(bytes_))
    tabla = next(s.table for s in prs.slides[1].shapes if s.has_table)
    fila = [c.text for c in tabla.rows[1].cells]
    assert "11.520.000" in fila[3]
